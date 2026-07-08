"""
第三阶段：从备件网 PDF 解析数据，填写 OMS 导出的 Excel。
"""
import re
import calendar
import logging
from datetime import datetime

from config import PRICE_VALIDITY_MONTHS

logger = logging.getLogger(__name__)


def _add_months(dt, months):
    month = dt.month - 1 + months
    year = dt.year + month // 12
    month = month % 12 + 1
    day = min(dt.day, calendar.monthrange(year, month)[1])
    return dt.replace(year=year, month=month, day=day)


def _norm_text(s):
    return re.sub(r"\s+", "", (s or "").lower())


def _col_index_by_header(headers, *keywords, exclude=()):
    """表头行中按关键词模糊找列索引（0-based）；exclude 排除含这些词的列。"""
    for i, h in enumerate(headers):
        hs = str(h or "")
        if any(ex in hs for ex in exclude):
            continue
        if all(kw in hs for kw in keywords):
            return i
    for i, h in enumerate(headers):
        hs = str(h or "")
        if any(ex in hs for ex in exclude):
            continue
        if any(kw in hs for kw in keywords):
            return i
    return None


def _find_oms_columns(headers, is_zhongshan=True):
    """按 OMS 导出表头精确定位列，避免「直发」误匹配 EXW/DDP 等列。"""
    col_exw = _col_index_by_header(
        headers, "EXW", exclude=("DDP", "成交", "包装", "交期", "有效")
    )
    if col_exw is None:
        col_exw = _col_index_by_header(headers, "出厂价", "未税", exclude=("DDP",))

    if is_zhongshan:
        # exclude「松江」仅用于表头定位：避免把「松江厂包装费」列当成中山的国内包装费列
        col_pack = _col_index_by_header(
            headers, "国内包装费", exclude=("松江",)
        )
        if col_pack is None:
            col_pack = _col_index_by_header(
                headers, "国内", "包装", exclude=("松江",)
            )
        col_freight = _col_index_by_header(headers, "直发工地", "运费")
        if col_freight is None:
            col_freight = _col_index_by_header(
                headers,
                "直发工地运费",
                exclude=("EXW", "DDP", "包装费", "包装尺寸", "单件重量"),
            )
    else:
        # 松江包装费列名也是「国内包装费」（与中山相同，但无需排除松江）
        col_pack = _col_index_by_header(headers, "国内包装费")
        if col_pack is None:
            col_pack = _col_index_by_header(headers, "国内", "包装")
        col_freight = None

    col_lead = _col_index_by_header(headers, "无库存", "交期")
    if col_lead is None:
        col_lead = _col_index_by_header(headers, "无库存", "天")

    col_valid_from = _col_index_by_header(
        headers, "价格有效从", exclude=("有效期至",)
    )
    col_valid_to = _col_index_by_header(
        headers, "价格有效期至", exclude=("有效从",)
    )
    if col_valid_to is None:
        col_valid_to = _col_index_by_header(
            headers, "有效期至", exclude=("有效从",)
        )

    col_desc = _col_index_by_header(headers, "物料", "描述", exclude=("单价",))
    if col_desc is None:
        col_desc = _col_index_by_header(headers, "描述", exclude=("单价",))

    col_mfg = _col_index_by_header(headers, "工厂料号")
    if col_mfg is None:
        col_mfg = _col_index_by_header(headers, "物料号")

    return {
        "desc": col_desc,
        "exw": col_exw,
        "pack": col_pack,
        "freight": col_freight,
        "lead": col_lead,
        "valid_from": col_valid_from,
        "valid_to": col_valid_to,
        "mfg": col_mfg,
    }


def _parse_delivery_days(pdf_text):
    """
    从 PDF 交货期段落解析天数（优先 Delivery/交货期 行，避免误匹配付款条款 30 days）。
    """
    lines = pdf_text.splitlines()
    delivery_lines = [
        line for line in lines if re.search(r"delivery|交货期", line, re.I)
    ]
    search_blocks = delivery_lines if delivery_lines else [pdf_text]

    patterns = [
        r"approx\.?\s*(\d+)\s*days?",
        r"约\s*(\d+)\s*天",
        r"(\d+)\s*days?\s*after\s+order",
        r"(\d+)\s*days?\s*after\s+order\s+confirmation",
        r"订单[^0-9]{0,20}(\d+)\s*天",
    ]
    for block in search_blocks:
        for pat in patterns:
            m = re.search(pat, block, re.I)
            if m:
                return int(m.group(1))

    m = re.search(r"交货期[^\d]{0,30}(\d+)", pdf_text)
    if m:
        return int(m.group(1))
    return None


def _parse_money(val):
    if val is None:
        return None
    s = str(val).replace(",", "").replace("￥", "").replace("¥", "").strip()
    m = re.search(r"[\d.]+", s)
    if not m:
        return None
    try:
        return float(m.group())
    except ValueError:
        return None


def _parse_oms_date(val):
    """解析 OMS 日期单元格（YYYYMMDD 整数/字符串、datetime 等）。"""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val
    s = re.sub(r"\D", "", str(val).strip())
    if len(s) >= 8:
        try:
            return datetime.strptime(s[:8], "%Y%m%d")
        except ValueError:
            pass
    s10 = str(val).strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s10, fmt)
        except ValueError:
            continue
    return None


def _format_oms_date(dt):
    """OMS 要求 YYYYMMDD 整数，如 20260819。"""
    return int(dt.strftime("%Y%m%d"))


def _extract_pdf_lines(pdf_path):
    """解析 PDF：交货期、物料行、包装费、运费。"""
    try:
        import pdfplumber
    except ImportError as e:
        raise ImportError(
            "缺少 pdfplumber，请执行: python -m pip install pdfplumber openpyxl"
        ) from e

    full_text = []
    line_items = []
    packaging = None
    freight = None

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            full_text.append(t)
            tables = page.extract_tables() or []
            for table in tables:
                if not table or len(table) < 2:
                    continue
                header = [str(c or "").strip() for c in table[0]]
                h_join = "".join(header)
                if "单价" not in h_join and "物料" not in h_join and "描述" not in h_join:
                    continue
                desc_i = _col_index_by_header(header, "物料") or _col_index_by_header(
                    header, "描述"
                )
                mfg_i = _col_index_by_header(header, "物料号") or _col_index_by_header(
                    header, "工厂料号"
                ) or _col_index_by_header(header, "料号")
                price_i = _col_index_by_header(header, "单价")
                if price_i is None:
                    price_i = _col_index_by_header(header, "EXW")
                for row in table[1:]:
                    if not row:
                        continue
                    cells = [str(c or "").strip() for c in row]
                    max_i = max(
                        filter(lambda x: x is not None, [desc_i, price_i, mfg_i]),
                        default=0,
                    )
                    if len(cells) <= max_i:
                        continue
                    desc = cells[desc_i] if desc_i is not None else ""
                    price = _parse_money(cells[price_i]) if price_i is not None else None
                    mfg_no = (cells[mfg_i] or "").strip() if mfg_i is not None else ""
                    if desc and price is not None:
                        line_items.append(
                            {"material_desc": desc, "unit_price": price, "mfg_no": mfg_no}
                        )

                for row in table[1:]:
                    if not row:
                        continue
                    row_text = " ".join(str(c or "") for c in row)
                    if packaging is None and re.search(
                        r"packing|包装费", row_text, re.I
                    ):
                        packaging = _parse_money(row_text)
                    if freight is None and re.search(
                        r"transportation|运费", row_text, re.I
                    ):
                        freight = _parse_money(row_text)

    text = "\n".join(full_text)
    days = _parse_delivery_days(text)

    if packaging is None:
        pm = re.search(
            r"PackingCost\s*/?\s*包装费[^\d]*([\d,]+\.?\d*)", text, re.I
        )
        if pm:
            packaging = _parse_money(pm.group(1))
        if packaging is None:
            pm = re.search(r"包装费[^\d￥¥]*[￥¥]?\s*([\d,]+\.?\d*)", text)
            if pm:
                packaging = _parse_money(pm.group(1))

    if freight is None:
        fm = re.search(
            r"Transportation\s+Fee[^\d]*([\d,]+\.?\d*)", text, re.I
        )
        if fm:
            freight = _parse_money(fm.group(1))
        if freight is None:
            fm = re.search(r"运费[^\d￥¥]*[￥¥]?\s*([\d,]+\.?\d*)", text)
            if fm:
                freight = _parse_money(fm.group(1))

    if packaging is None:
        pm2 = re.search(r"packaging[^\d]*([\d,]+\.?\d*)", text, re.I)
        if pm2:
            packaging = _parse_money(pm2.group(1))

    return {
        "delivery_days": days,
        "line_items": line_items,
        "packaging": packaging,
        "freight": freight,
        "full_text": text,
    }


def _find_header_row(ws, max_scan=30):
    keywords = ["物料", "EXW", "出厂价", "交期", "包装"]
    for r in range(1, min(max_scan, ws.max_row) + 1):
        row_vals = [
            str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1)
        ]
        joined = "".join(row_vals)
        if any(k in joined for k in keywords):
            return r, row_vals
    return None, []


def _count_material_rows(ws, header_row, col_desc):
    """统计 Excel 中有物料描述的数据行数（类数），用于包装费/运费按行均摊。"""
    if col_desc is None:
        return 0
    count = 0
    for r in range(header_row + 1, ws.max_row + 1):
        excel_desc = ws.cell(r, col_desc + 1).value
        if excel_desc and str(excel_desc).strip():
            count += 1
    return count


def _match_row_desc(excel_desc, pdf_desc):
    a, b = _norm_text(excel_desc), _norm_text(pdf_desc)
    if not a or not b:
        return False
    if a == b or a in b or b in a:
        return True
    key_a = a[: min(12, len(a))]
    key_b = b[: min(12, len(b))]
    return key_a in b or key_b in a


def fill_excel_from_pdf(
    pdf_path,
    xlsx_path,
    factory_type="中山",
    project_split_count=1,
):
    """
    按业务规则填写 Excel 并覆盖保存。
    包装费/运费：PDF 总额 ÷ Excel 物料行数（类数），每行写入相同均摊值。
    中山填国内包装费+直发运费；松江仅填松江厂包装费列，不写运费。
    返回 dict 含填写摘要。
    """
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ImportError(
            "缺少 openpyxl，请执行: python -m pip install openpyxl pdfplumber"
        ) from e

    is_zhongshan = "中山" in (factory_type or "")
    if project_split_count not in (None, 1):
        logger.warning(
            f"  project_split_count={project_split_count} 已忽略；"
            "包装费/运费按 Excel 物料行数均摊"
        )

    pdf_data = _extract_pdf_lines(pdf_path)
    if not pdf_data["line_items"]:
        logger.warning("  PDF 未解析到物料单价行，将仅填交期/有效期/包装运费")

    wb = load_workbook(xlsx_path)
    ws = wb.active
    header_row, headers = _find_header_row(ws)
    if not header_row:
        raise RuntimeError("Excel 中未找到表头行")

    cols = _find_oms_columns(headers, is_zhongshan=is_zhongshan)
    col_desc = cols["desc"]
    col_exw = cols["exw"]
    col_pack = cols["pack"]
    col_freight = cols["freight"]
    col_lead = cols["lead"]
    col_valid_from = cols["valid_from"]
    col_valid_to = cols["valid_to"]

    material_line_count = max(1, _count_material_rows(ws, header_row, col_desc))

    logger.info(
        f"  列映射: 交期={col_lead} 包装={col_pack} 运费={col_freight} "
        f"EXW={col_exw} 有效从={col_valid_from} 有效至={col_valid_to} "
        f"物料行数={material_line_count}"
    )

    packaging_val = pdf_data["packaging"]
    freight_val = pdf_data["freight"]
    if packaging_val is not None:
        packaging_val = packaging_val / material_line_count
    if freight_val is not None and is_zhongshan:
        freight_val = freight_val / material_line_count

    filled_prices = 0
    data_start = header_row + 1
    for r in range(data_start, ws.max_row + 1):
        excel_desc = ws.cell(r, col_desc + 1).value if col_desc is not None else None
        if not excel_desc or not str(excel_desc).strip():
            continue

        if col_lead is not None and pdf_data["delivery_days"] is not None:
            ws.cell(r, col_lead + 1).value = pdf_data["delivery_days"]

        if col_pack is not None and packaging_val is not None:
            ws.cell(r, col_pack + 1).value = packaging_val

        if col_freight is not None and is_zhongshan and freight_val is not None:
            ws.cell(r, col_freight + 1).value = freight_val

        if col_exw is not None:
            col_mfg = cols.get("mfg")
            excel_mfg = (
                str(ws.cell(r, col_mfg + 1).value or "").strip()
                if col_mfg is not None else ""
            )

            matched = None
            # 优先：按工厂料号精确匹配（PDF 中为「物料号」列）
            if excel_mfg:
                for item in pdf_data["line_items"]:
                    if (item.get("mfg_no") or "").strip() == excel_mfg:
                        matched = item
                        break
            # 降级：按物料描述匹配（工厂料号为空或匹配失败时）
            if matched is None:
                for item in pdf_data["line_items"]:
                    if _match_row_desc(excel_desc, item["material_desc"]):
                        matched = item
                        break

            if matched:
                ws.cell(r, col_exw + 1).value = matched["unit_price"]
                filled_prices += 1
            else:
                logger.error(
                    f"  物料无法匹配 PDF 单价: Excel={str(excel_desc)[:40]}"
                    f" 工厂料号={excel_mfg or '—'}"
                )

        # 仅改「价格有效期至」=「价格有效从」+3 个月；不修改「价格有效从」
        if col_valid_from is not None and col_valid_to is not None:
            vf = ws.cell(r, col_valid_from + 1).value
            dt_from = _parse_oms_date(vf)
            if dt_from:
                dt_to = _add_months(dt_from, PRICE_VALIDITY_MONTHS)
                ws.cell(r, col_valid_to + 1).value = _format_oms_date(dt_to)
            elif vf:
                logger.warning(f"  无法解析价格有效从: {vf!r}，跳过有效期至")

    try:
        wb.save(xlsx_path)
    except PermissionError as e:
        raise PermissionError(
            f"无法保存 Excel（文件可能正被 Excel 打开，请先关闭）: {xlsx_path}"
        ) from e

    summary = {
        "delivery_days": pdf_data["delivery_days"],
        "filled_prices": filled_prices,
        "pdf_lines": len(pdf_data["line_items"]),
        "packaging": packaging_val,
        "freight": freight_val if is_zhongshan else None,
        "material_line_count": material_line_count,
        "columns": cols,
    }
    logger.info(f"  Excel 填表完成: {summary}")
    return summary

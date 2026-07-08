"""
第三阶段填表逻辑单元测试（无需浏览器）。
运行: python test_quote_fill.py
"""
import os
import sys
import unittest
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from modules.quote_fill import (
    _parse_delivery_days,
    _match_row_desc,
    _add_months,
    _parse_money,
    _norm_text,
    _parse_oms_date,
    _format_oms_date,
    _find_oms_columns,
    _count_material_rows,
)


class TestQuoteFillHelpers(unittest.TestCase):
    def test_parse_delivery_days(self):
        self.assertEqual(_parse_delivery_days("Delivery: 14 days after order"), 14)
        self.assertEqual(
            _parse_delivery_days(
                "Payment T/T 30 days after B/L\n"
                "Delivery/ 交货期: Approx. 14 days after order confirmation."
            ),
            14,
        )
        self.assertEqual(_parse_delivery_days("交货期：21天"), 21)
        self.assertIsNone(_parse_delivery_days("no date here"))

    def test_parse_oms_date_and_valid_to(self):
        dt = _parse_oms_date(20260519)
        self.assertEqual((dt.year, dt.month, dt.day), (2026, 5, 19))
        dt_to = _add_months(dt, 3)
        self.assertEqual(_format_oms_date(dt_to), 20260819)

    def test_find_oms_columns_freight_not_exw(self):
        headers = [""] * 23
        headers[12] = "*EXW出厂价未税\n(若整车直发安装时需在物料描述中填写)"
        headers[20] = "国内包装费\n(中山厂整车直发时需要填写)"
        headers[22] = "直发工地运费\n(中山整车直发时需要填写)"
        cols = _find_oms_columns(headers, is_zhongshan=True)
        self.assertEqual(cols["freight"], 22)
        self.assertEqual(cols["pack"], 20)

    def test_find_oms_columns_songjiang_pack_only(self):
        headers = [""] * 20
        headers[15] = "松江厂包装费"
        cols = _find_oms_columns(headers, is_zhongshan=False)
        self.assertEqual(cols["pack"], 15)
        self.assertIsNone(cols["freight"])

    def test_match_row_desc(self):
        self.assertTrue(_match_row_desc("油压缓冲器 tkOB10B", "油压缓冲器 TKOB10B PC"))
        self.assertTrue(_match_row_desc("缓冲器", "油压缓冲器 tkOB10B"))
        self.assertFalse(_match_row_desc("abc", "xyz completely different"))

    def test_add_months(self):
        d = datetime(2026, 1, 31)
        r = _add_months(d, 3)
        self.assertEqual((r.year, r.month), (2026, 4))
        r2 = _add_months(datetime(2026, 11, 15), 3)
        self.assertEqual((r2.year, r2.month), (2027, 2))

    def test_parse_money(self):
        self.assertEqual(_parse_money("￥1,168.80"), 1168.80)
        self.assertEqual(_parse_money("389.60"), 389.60)
        self.assertIsNone(_parse_money(""))

    def test_norm_text(self):
        self.assertEqual(_norm_text("A B\tC"), "abc")

    def test_count_material_rows(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl not installed")
        wb = Workbook()
        ws = wb.active
        ws.cell(1, 1).value = "物料描述"
        ws.cell(2, 1).value = "缓冲器A"
        ws.cell(3, 1).value = "缓冲器B"
        ws.cell(4, 1).value = ""
        ws.cell(5, 1).value = "缓冲器C"
        self.assertEqual(_count_material_rows(ws, 1, 0), 3)


if __name__ == "__main__":
    unittest.main()
